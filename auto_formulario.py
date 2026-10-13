#!/usr/bin/env python3
"""
Auto-respondedor de formularios web con Playwright (sync API) + Claude.

Flujo:
  1. Abre un navegador visible; tú navegas/inicias sesión a mano.
  2. El script espera en la terminal hasta que presiones ENTER.
  3. Escanea la página, clasifica cada pregunta (abierta / única / múltiple),
     consulta a Claude y rellena las respuestas.

Instalación:
    pip install playwright anthropic python-pptx pypdf   # (python-pptx/pypdf solo si usas --material)
    playwright install chromium
    export ANTHROPIC_API_KEY="sk-ant-..."      # (Windows PowerShell: $env:ANTHROPIC_API_KEY="...")

Uso:
    python auto_formulario.py --url "https://tu-sitio.com/formulario"
    python auto_formulario.py --backend cli     # usa el CLI `claude -p` en vez de la API
    python auto_formulario.py --backend cli --material ./presentaciones   # responde usando tus .pptx/.pdf/.txt/.md
    python auto_formulario.py --demo            # prueba con un formulario local de ejemplo
    python auto_formulario.py --demo-material --backend cli --material ./material   # verifica que usa TU material

Modelos: por defecto Sonnet para todo. Opciones: --model X (un solo modelo), o enrutar por
dificultad con --model-facil claude-haiku-5-5 --model-dificil claude-sonnet-5-5 (o claude-opus-5-5).
Al final se imprime el consumo de tokens de la sesión.

Usa esta herramienta solo en formularios/evaluaciones donde esté permitido
(prácticas, encuestas, formularios propios, etc.).
"""

import argparse
import base64
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from urllib.parse import quote

from playwright.sync_api import Locator, Page, sync_playwright

# ════════════════════════════════════════════════════════════════════════════
# CONFIGURACIÓN — PERSONALIZA ESTOS SELECTORES SEGÚN EL SITIO REAL
# ════════════════════════════════════════════════════════════════════════════

# Contenedor de CADA pregunta (el bloque que agrupa enunciado + campos/opciones).
# Ej.: ".pregunta", ".question", "div.que", "fieldset", "[data-qa='question']"
SELECTOR_PREGUNTA = ".pregunta"

# Elemento (dentro del contenedor) con el texto del enunciado.
# Ej.: ".enunciado", ".qtext", "legend", "h3". Si no existe, se usa el texto
# completo del contenedor menos el de las opciones.
SELECTOR_ENUNCIADO = ".enunciado"

# Contenedor de CADA opción (radio/checkbox + su texto). Es opcional: si no
# coincide nada, se usan directamente los <input> y se deduce su etiqueta.
# Ej.: ".opcion", ".answer label", "li.choice"
SELECTOR_OPCION = ".opcion"

# Campos de las posibles respuestas (normalmente no hace falta cambiarlos).
# Editores de texto enriquecido (p. ej. TinyMCE en Canvas) viven en un <iframe> o en un [contenteditable].
SELECTOR_EDITOR_IFRAME = "iframe[id$='_ifr'], iframe.tox-edit-area__iframe"
SELECTOR_ABIERTA = ("textarea, input[type='text'], input[type='number'], input:not([type]), "
                    "[contenteditable='true'], " + SELECTOR_EDITOR_IFRAME)
SELECTOR_DESPLEGABLE = "select"  # listas desplegables (p. ej. Verdadero/Falso por fila)
SELECTOR_RADIO = "input[type='radio'], [role='radio']"
SELECTOR_CHECKBOX = "input[type='checkbox'], [role='checkbox']"

# Botón de enviar. Déjalo en None para NO enviar (recomendado: revisa tú).
SELECTOR_BOTON_ENVIAR = None  # Ej.: "button[type='submit']"

# Por defecto se usa Sonnet para TODO (más seguro para acertar). Enrutamiento opcional: si
# MODELO_FACIL != MODELO_DIFICIL (ej. --model-facil claude-haiku-5-5), una consulta barata
# (sin el material) clasifica cada pregunta como FÁCIL o DIFÍCIL y elige el modelo.
MODELO_FACIL = os.getenv("CLAUDE_MODEL_FACIL", "claude-sonnet-5-5")
MODELO_DIFICIL = os.getenv("CLAUDE_MODEL_DIFICIL", "claude-sonnet-5-5")  # o "claude-opus-5-5"
ENRUTAR = False  # se activa solo si los dos modelos son distintos (ver main)

# Si el enunciado contiene estas palabras se considera DIFÍCIL sin gastar la consulta de clasificación.
PATRON_DIFICIL = re.compile(
    r"calcul|resuelv|resuelve|determina|demuestr|deduc|deriv|despej|desarroll|paso a paso|"
    r"procedimiento|encuentra el valor|obt[eé]n", re.I)
TIMEOUT_ACCION_MS = 5000

SYSTEM_PROMPT = (
    "Eres un asistente experto que responde preguntas de formularios. "
    "Sigue estrictamente el formato de salida pedido, sin explicaciones, "
    "sin comillas ni texto adicional."
)


# ════════════════════════════════════════════════════════════════════════════
# a) ESPERAR EL ACTIVADOR DEL USUARIO
# ════════════════════════════════════════════════════════════════════════════
def esperar_activador() -> None:
    """Bloquea hasta que el usuario presione ENTER en la terminal."""
    print("\n" + "=" * 60)
    print(" Navega en el navegador hasta la página con las preguntas.")
    print(" Cuando estés en la pantalla correcta, vuelve aquí y presiona ENTER.")
    print("=" * 60)
    try:
        input(">>> Presiona ENTER para iniciar el escaneo... ")
    except EOFError:
        print("(No hay terminal interactiva; continuando.)")


# ════════════════════════════════════════════════════════════════════════════
# MATERIAL DE CONSULTA (presentaciones, PDFs, notas) — opcional con --material
# ════════════════════════════════════════════════════════════════════════════
CONTEXTO_MATERIAL = ""  # se llena en main() si usas --material
MAX_CHARS_MATERIAL = 400_000  # ~100k tokens; evita desbordar el contexto de Claude


def _texto_pptx(ruta: Path) -> str:
    from pptx import Presentation  # pip install python-pptx

    def textos(shapes):
        for sh in shapes:
            if sh.shape_type == 6:  # grupo de formas
                yield from textos(sh.shapes)
            if sh.has_text_frame and sh.text_frame.text.strip():
                yield sh.text_frame.text.strip()
            if getattr(sh, "has_table", False) and sh.has_table:
                for fila in sh.table.rows:
                    yield " | ".join(c.text.strip() for c in fila.cells)

    partes = []
    for n, sl in enumerate(Presentation(str(ruta)).slides, 1):
        lineas = list(textos(sl.shapes))
        if sl.has_notes_slide and sl.notes_slide.notes_text_frame.text.strip():
            lineas.append("Notas: " + sl.notes_slide.notes_text_frame.text.strip())
        partes.append(f"[Diapositiva {n}]\n" + "\n".join(lineas))
    return "\n\n".join(partes)


def _texto_pdf(ruta: Path) -> str:
    from pypdf import PdfReader  # pip install pypdf

    return "\n\n".join(f"[Página {n}]\n{(pg.extract_text() or '').strip()}"
                        for n, pg in enumerate(PdfReader(str(ruta)).pages, 1))


def cargar_material(ruta: str) -> str:
    """Lee un archivo o carpeta (.pptx, .pdf, .txt, .md) y devuelve todo su texto."""
    base = Path(ruta).expanduser()
    if not base.exists():
        sys.exit(f"--material: no existe la ruta '{ruta}'")
    archivos = [base] if base.is_file() else sorted(p for p in base.rglob("*") if p.is_file())
    bloques = []
    for a in archivos:
        ext = a.suffix.lower()
        try:
            if ext == ".pptx":
                txt = _texto_pptx(a)
            elif ext == ".pdf":
                txt = _texto_pdf(a)
            elif ext in (".txt", ".md"):
                txt = a.read_text(encoding="utf-8", errors="ignore")
            else:
                continue  # .ppt antiguo, imágenes, etc.: se ignoran
        except Exception as e:
            print(f"  (no se pudo leer {a.name}: {type(e).__name__}: {e})")
            continue
        if txt.strip():
            print(f"  Material cargado: {a.name} ({len(txt):,} caracteres)")
            bloques.append(f"=== ARCHIVO: {a.name} ===\n{txt.strip()}")
    texto = "\n\n".join(bloques)
    if not texto:
        sys.exit("--material: no se encontró texto legible (.pptx/.pdf/.txt/.md).")
    if len(texto) > MAX_CHARS_MATERIAL:
        print(f"  AVISO: material muy largo; se recorta a {MAX_CHARS_MATERIAL:,} caracteres.")
        texto = texto[:MAX_CHARS_MATERIAL]
    return texto


def construir_system() -> str:
    if not CONTEXTO_MATERIAL:
        return SYSTEM_PROMPT
    return (
        f"{SYSTEM_PROMPT}\n\nTu fuente principal es el MATERIAL DE REFERENCIA de abajo. "
        "Basa tus respuestas en él; si la respuesta no aparece ahí, usa tu conocimiento general.\n\n"
        f"<material_de_referencia>\n{CONTEXTO_MATERIAL}\n</material_de_referencia>"
    )


# ════════════════════════════════════════════════════════════════════════════
# b) CONSULTAR A CLAUDE
# ════════════════════════════════════════════════════════════════════════════
CONSUMO = {"llamadas": 0, "nuevos": 0, "cache": 0, "salida": 0, "usd": 0.0}
_ARCHIVOS_SYSTEM: dict[str, str] = {}


def _archivo_system(texto: str) -> str:
    """Guarda el system prompt (con el material) en un archivo temporal, una sola vez."""
    if texto not in _ARCHIVOS_SYSTEM:
        ruta = Path(tempfile.mkdtemp()) / "system.txt"
        ruta.write_text(texto, encoding="utf-8")
        _ARCHIVOS_SYSTEM[texto] = str(ruta)
    return _ARCHIVOS_SYSTEM[texto]


def _registrar_consumo(uso: dict, usd: float | None) -> None:
    CONSUMO["llamadas"] += 1
    CONSUMO["nuevos"] += (uso.get("input_tokens") or 0) + (uso.get("cache_creation_input_tokens") or 0)
    CONSUMO["cache"] += uso.get("cache_read_input_tokens") or 0
    CONSUMO["salida"] += uso.get("output_tokens") or 0
    CONSUMO["usd"] += usd or 0.0


def imprimir_consumo() -> None:
    c = CONSUMO
    if not c["llamadas"]:
        return
    print("\n" + "-" * 60)
    print(f" Consumo de esta sesión: {c['llamadas']} llamadas a Claude")
    print(f"   tokens de entrada nuevos: {c['nuevos']:,}   (caché leído: {c['cache']:,})")
    print(f"   tokens de salida: {c['salida']:,}   | costo equivalente en API: ${c['usd']:.2f}")
    print("   (Tu límite de suscripción no se mide en dólares; revisa /usage en `claude` antes y después.)")
    print("-" * 60)


def _bloques_imagen(imagenes: list[bytes]) -> list[dict]:
    return [{"type": "image", "source": {"type": "base64", "media_type": "image/png",
                                         "data": base64.b64encode(img).decode()}} for img in imagenes]


def consultar_claude(prompt: str, backend: str = "api", modelo: str | None = None,
                     usar_material: bool = True, max_tokens: int = 1024,
                     imagenes: list[bytes] | None = None) -> str:
    """Envía `prompt` (y opcionalmente imágenes PNG) a Claude y devuelve el texto de la respuesta."""
    modelo = modelo or MODELO_FACIL
    system = construir_system() if usar_material else SYSTEM_PROMPT
    imagenes = imagenes or []

    if backend == "api":
        import anthropic  # import perezoso: solo si se usa

        cliente = anthropic.Anthropic()  # lee ANTHROPIC_API_KEY del entorno
        resp = cliente.messages.create(
            model=modelo,
            max_tokens=max_tokens,
            system=system,
            messages=[{"role": "user", "content": _bloques_imagen(imagenes) + [{"type": "text", "text": prompt}]}],
        )
        return "".join(b.text for b in resp.content if b.type == "text").strip()

    if backend == "cli":
        # Requiere tener instalado y autenticado el CLI de Claude Code.
        # El material va como system prompt (archivo): así Claude puede reutilizarlo desde su caché
        # entre preguntas (≈20x menos consumo que mandarlo dentro de cada pregunta).
        cmd = ["claude", "-p", "--tools", "", "--model", modelo,
               "--append-system-prompt-file", _archivo_system(system)]
        if imagenes:  # con imágenes la entrada es JSON por stdin (stream-json)
            cmd += ["--input-format", "stream-json", "--output-format", "stream-json", "--verbose"]
            entrada = json.dumps({"type": "user", "message": {"role": "user",
                                  "content": _bloques_imagen(imagenes) + [{"type": "text", "text": prompt}]}}) + "\n"
        else:
            cmd += ["--output-format", "json"]
            entrada = prompt
        error = ""
        for intento in (1, 2):  # un reintento si el CLI se cuelga o devuelve vacío
            try:
                res = subprocess.run(cmd, input=entrada, capture_output=True, text=True,
                                     encoding="utf-8", timeout=240)
                if res.returncode == 0 and res.stdout.strip():
                    texto = _leer_resultado_cli(res.stdout, stream=bool(imagenes))
                    if texto:
                        return texto
                error = res.stderr.strip() or "respuesta vacía"
            except subprocess.TimeoutExpired:
                error = "tiempo de espera agotado"
            if intento == 1:
                print(f"     (reintentando: {error[:80]})")
        raise RuntimeError(f"claude CLI falló: {error}")

    if backend == "mock":
        # Respuestas falsas SOLO para probar el flujo sin gastar API.
        if "SOLO la palabra FACIL o DIFICIL" in prompt:
            return "FACIL"
        if "una letra por hueco" in prompt:
            return '["B", "B", "A", "B", "B", "B", "A", "B"]'
        if "arreglo JSON" in prompt:
            return '["A", "C"]'
        if "SOLO la letra" in prompt:
            return "B"
        return "Respuesta de prueba (mock)."

    raise ValueError(f"backend desconocido: {backend}")


def _leer_resultado_cli(stdout: str, stream: bool) -> str:
    """Extrae el texto final (y registra consumo) de la salida JSON / stream-json del CLI."""
    datos = None
    if stream:
        for linea in stdout.splitlines():
            try:
                ev = json.loads(linea)
            except json.JSONDecodeError:
                continue
            if ev.get("type") == "result":
                datos = ev
    else:
        try:
            datos = json.loads(stdout)
        except json.JSONDecodeError:
            return stdout.strip()
    if not datos:
        return ""
    _registrar_consumo(datos.get("usage") or {}, datos.get("total_cost_usd"))
    return (datos.get("result") or "").strip()


def clasificar_dificultad(enunciado: str, opciones: list[str], backend: str) -> bool:
    """True si la pregunta es DIFÍCIL (cálculo/desarrollo). Si algo falla, asume difícil (más seguro)."""
    if not ENRUTAR:
        return False
    if PATRON_DIFICIL.search(enunciado) or len(enunciado) > 700:
        return True
    prompt = (
        "Clasifica la dificultad de esta pregunta de examen. Responde SOLO la palabra FACIL o DIFICIL.\n"
        "FACIL: basta recordar o ubicar un dato, definición, concepto o lista.\n"
        "DIFICIL: requiere cálculos, desarrollo matemático o algebraico de varios pasos, aplicar un "
        "procedimiento o razonar con varias condiciones.\n\n"
        f"Pregunta: {enunciado}\n"
        + (f"Opciones:\n{_lista_letras(opciones)}\n" if opciones else "")
    )
    try:
        r = consultar_claude(prompt, backend, modelo=MODELO_FACIL, usar_material=False, max_tokens=10)
        return "DIFICIL" in r.upper().replace("Í", "I")
    except Exception:
        return True


def _nota_imagenes(imagenes: list[bytes] | None) -> str:
    if not imagenes:
        return ""
    return (f"\n(Se adjuntan {len(imagenes)} imagen(es) de la pregunta, en orden de aparición "
            "([IMAGEN n] en el texto): úsalas, pueden contener datos, gráficas o fórmulas.)\n")


_REGLAS_ABIERTA = (
    "Si tiene incisos (a, b, c…), responde cada uno etiquetado con su letra, p. ej. 'a) ...'. "
    "Escribe las fórmulas en texto plano (ej. sqrt(x), x^2, a/b, derivada parcial dU/dX_A), sin markdown ni LaTeX. "
    "Devuelve SOLO la respuesta."
)


def _respuesta_final(texto: str) -> str:
    """Toma lo que sigue a la última 'Respuesta final:' (o la última línea) para campos de una línea."""
    m = re.findall(r"respuesta final\s*:\s*(.+)", texto, re.I)
    res = m[-1] if m else (texto.strip().splitlines() or [""])[-1]
    return res.strip().strip("*").strip().rstrip(".")


def pedir_abierta(enunciado: str, backend: str, modelo: str, dificil: bool = False,
                  imagenes: list[bytes] | None = None, corta: bool = False) -> str:
    base = f"{enunciado}{_nota_imagenes(imagenes)}"
    if corta:  # el campo es de una sola línea: se razona, pero solo se escribe el resultado final
        prompt = (
            f"Resuelve:\n\n{base}\n\n"
            "Razona paso a paso de forma compacta y termina con una última línea 'Respuesta final: ...' que "
            "contenga SOLO el valor o la respuesta final (un número, expresión corta o palabra), sin explicación. "
            "Texto plano, sin markdown ni LaTeX."
        )
        return _respuesta_final(consultar_claude(prompt, backend, modelo, max_tokens=2048, imagenes=imagenes))
    if dificil:
        prompt = (
            f"Resuelve este ejercicio:\n\n{base}\n\n"
            "Desarrolla el procedimiento paso a paso de forma ordenada y compacta (fórmulas y cálculos "
            f"clave, sin explicaciones largas) y termina con una línea 'Respuesta final: ...'. {_REGLAS_ABIERTA}"
        )
    else:
        prompt = (
            f"Responde a:\n\n{base}\n\n"
            "Si es una pregunta de datos, definiciones o conceptos, responde de forma concisa y directa "
            "(máximo 2-3 frases). Si pide cálculos, un desarrollo o una demostración, resuélvelo paso a paso "
            f"de forma ordenada y compacta y termina con una línea 'Respuesta final: ...'. {_REGLAS_ABIERTA}"
        )
    return consultar_claude(prompt, backend, modelo, max_tokens=2048, imagenes=imagenes)


def _letra(i: int) -> str:
    return chr(ord("A") + i) if i < 26 else str(i + 1)


def _lista_letras(opciones: list[str]) -> str:
    # Letras (A, B, C...) en vez de números: así no se confunden con opciones numéricas.
    return "\n".join(f"{_letra(i)}) {t}" for i, t in enumerate(opciones))


def _letras_a_indices(texto: str, opciones: list[str]) -> list[int]:
    """Convierte la respuesta de Claude (letras, o el texto de las opciones) en índices 0-based."""
    n = len(opciones)
    # Se quitan comillas/corchetes y se buscan letras sueltas: A, B, "C", [A, C], A) ...
    limpio = texto.replace("\n", " ")
    for patron in (r"(?<![A-Za-zÁ-ú0-9])([A-Z])(?![A-Za-zÁ-ú0-9])",   # primero mayúsculas
                   r"(?<![A-Za-zÁ-ú0-9])([a-z])(?![A-Za-zÁ-ú0-9])"):  # luego minúsculas ("b.")
        idx = sorted({ord(l.upper()) - 65 for l in re.findall(patron, limpio)
                      if 0 <= ord(l.upper()) - 65 < n})
        if idx:
            return idx
    # Respaldo: Claude devolvió el texto exacto de las opciones.
    partes = {x.strip().strip("\"'[]").lower() for x in re.split(r"[,\n;]", texto)}
    return [i for i, op in enumerate(opciones) if op.strip().lower() in partes]


def pedir_unica(enunciado: str, opciones: list[str], backend: str, modelo: str,
                imagenes: list[bytes] | None = None) -> int | None:
    """Devuelve el índice (0-based) de la opción elegida, o None si no se pudo."""
    prompt = (
        f"Pregunta:\n{enunciado}\n{_nota_imagenes(imagenes)}\nOpciones:\n{_lista_letras(opciones)}\n\n"
        "Elige la ÚNICA opción correcta. Devuelve SOLO la letra de la opción "
        "(por ejemplo: B). Nada más."
    )
    idx = _letras_a_indices(consultar_claude(prompt, backend, modelo, imagenes=imagenes), opciones)
    return idx[0] if idx else None


def pedir_multiple(enunciado: str, opciones: list[str], backend: str, modelo: str,
                   imagenes: list[bytes] | None = None) -> list[int]:
    """Devuelve la lista de índices (0-based) de las opciones correctas."""
    prompt = (
        f"Pregunta:\n{enunciado}\n{_nota_imagenes(imagenes)}\nOpciones:\n{_lista_letras(opciones)}\n\n"
        "Selecciona TODAS las opciones correctas (una o más). Devuelve SOLO un "
        'arreglo JSON con las letras, por ejemplo: ["A", "C"]. Nada más.'
    )
    return _letras_a_indices(consultar_claude(prompt, backend, modelo, imagenes=imagenes), opciones)


def _letras_en_orden(texto: str, n: int) -> list[str]:
    """Letras de la respuesta en el orden en que aparecen (se pueden repetir), máximo n."""
    letras: list[str] = []
    try:
        m = re.search(r"\[.*\]", texto, re.S)
        letras = [str(x).strip().upper()[:1] for x in (json.loads(m.group()) if m else [])]
    except (json.JSONDecodeError, TypeError):
        letras = []
    if not letras:
        letras = [l.upper() for l in re.findall(r"(?<![A-Za-zÁ-ú0-9])([A-Za-z])(?![A-Za-zÁ-ú0-9])", texto)]
    return letras[:n]


def pedir_desplegables(enunciado: str, huecos: list[dict], backend: str, modelo: str,
                       imagenes: list[bytes] | None = None) -> list[int | None]:
    """Un índice de opción (0-based, dentro de las opciones reales de cada hueco) por cada <select>."""
    listas = [h["textos"] for h in huecos]
    iguales = all(l == listas[0] for l in listas)
    if iguales:
        opciones = f"Opciones (iguales en todos los huecos):\n{_lista_letras(listas[0])}"
    else:
        opciones = "\n\n".join(f"Opciones del hueco [[{i}]]:\n{_lista_letras(l)}" for i, l in enumerate(listas, 1))
    contextos = "\n".join(f"[[{i}]] corresponde a: {h['contexto']}" for i, h in enumerate(huecos, 1) if h["contexto"])
    prompt = (
        "Pregunta con listas desplegables. Cada hueco aparece como [[1]], [[2]]… junto al texto al que se refiere:\n\n"
        f"{enunciado}\n{_nota_imagenes(imagenes)}\n"
        f"{contextos + chr(10) if contextos else ''}\n{opciones}\n\n"
        f"Elige la opción correcta de CADA uno de los {len(huecos)} huecos. Devuelve SOLO un arreglo JSON con "
        'una letra por hueco, en orden, por ejemplo: ["B", "A", "A"]. Nada más.'
    )
    letras = _letras_en_orden(consultar_claude(prompt, backend, modelo, imagenes=imagenes), len(huecos))
    out: list[int | None] = []
    for i, h in enumerate(huecos):
        k = ord(letras[i]) - 65 if i < len(letras) else -1
        out.append(k if 0 <= k < len(h["textos"]) else None)
    return out


# ════════════════════════════════════════════════════════════════════════════
# c) DETECTAR EL TIPO DE PREGUNTA E INTERACTUAR CON LA PÁGINA
# ════════════════════════════════════════════════════════════════════════════
# Función JS que lee el texto de un bloque de forma fiable:
#  - fórmulas (KaTeX, MathJax, <math>, imágenes de ecuación de Canvas) -> su código LaTeX original, sin duplicados
#    (copiar/pegar lo renderizado puede perder signos o duplicar símbolos);
#  - imágenes reales (gráficas, diagramas) -> marcador [IMAGEN n] y se devuelven sus ids para capturarlas;
#  - <select> -> "____", o [[n]] si marcarSelects es true.
_JS_TEXTO_FN = r"""function (el, marcarSelects) {
  window.__afImg = window.__afImg || 0;
  const MATH = '.katex, mjx-container, .MathJax, .MathJax_Preview, .MathJax_SVG, .MathJax_Display, .MathJax_SVG_Display, math';
  const esEcuacion = n => n.matches('img.equation_image, img[data-equation-content]') || /^latex/i.test(n.getAttribute('alt') || '');
  const ids = [];
  el.querySelectorAll('img, canvas, svg').forEach(n => {
    if (n.closest(MATH)) return;                                   // gráficos internos de fórmulas
    if (n.tagName.toLowerCase() === 'img' && esEcuacion(n)) return; // imágenes de ecuación
    if (n.parentElement && n.parentElement.closest('svg')) return;  // svg anidado
    const r = n.getBoundingClientRect();
    if (r.width < 40 || r.height < 40) return;                      // iconos
    let id = n.getAttribute('data-af-img');
    if (id === null) { id = String(window.__afImg++); n.setAttribute('data-af-img', id); }
    ids.push(id);
  });
  const clon = el.cloneNode(true);
  const T = s => document.createTextNode(' ' + String(s).replace(/\s+/g, ' ').trim() + ' ');
  const tex = n => { const a = n.querySelector('annotation[encoding="application/x-tex"]') || n.querySelector('annotation'); return a ? a.textContent : null; };
  clon.querySelectorAll('.katex').forEach(k => { const t = tex(k); k.replaceWith(T(t !== null ? t : k.textContent)); });
  clon.querySelectorAll('mjx-container').forEach(m => {
    const t = tex(m), mm = m.querySelector('mjx-assistive-mml');
    m.replaceWith(T(t !== null ? t : (mm ? mm.textContent : (m.getAttribute('aria-label') || m.textContent))));
  });
  const scripts = clon.querySelectorAll('script[type^="math/tex"]');
  scripts.forEach(sc => sc.replaceWith(T(sc.textContent)));
  if (scripts.length) clon.querySelectorAll('.MathJax_Preview, .MathJax, .MathJax_Display, .MathJax_SVG, .MathJax_SVG_Display').forEach(m => m.remove());
  clon.querySelectorAll('img').forEach(i => {
    if (esEcuacion(i)) i.replaceWith(T(i.getAttribute('data-equation-content') || (i.getAttribute('alt') || '').replace(/^latex:\s*/i, '') || i.getAttribute('title') || ''));
  });
  clon.querySelectorAll('math').forEach(m => { const t = tex(m); m.replaceWith(T(t !== null ? t : m.textContent)); });
  let k = 0;
  clon.querySelectorAll('select').forEach(sl => sl.replaceWith(T(marcarSelects ? '[[' + (++k) + ']]' : '____')));
  clon.querySelectorAll('[data-af-img]').forEach((n, i) => n.replaceWith(T('[IMAGEN ' + (i + 1) + ']')));
  clon.querySelectorAll('script, style, noscript, textarea, iframe').forEach(n => n.remove());
  clon.style.cssText = 'position:absolute;left:-99999px;top:0;width:' + Math.max(el.getBoundingClientRect().width, 300) + 'px';
  document.body.appendChild(clon);
  const texto = clon.innerText;
  clon.remove();
  return { texto: texto, ids: ids };
}"""
_JS_CONTENIDO = "(el, marcar) => (" + _JS_TEXTO_FN + ")(el, marcar)"

# Texto asociado a un radio/checkbox: <label for>, label padre, aria-label, o el contenedor más
# cercano que no incluya otros inputs (con fórmulas convertidas a LaTeX).
_JS_ETIQUETA = r"""el => {
  const fn = """ + _JS_TEXTO_FN + r""";
  let t = (el.labels && el.labels.length) ? el.labels[0] : el.closest('label');
  if (!t) { const aria = el.getAttribute('aria-label'); if (aria) return aria; }
  if (!t) {
    let n = el;
    for (let i = 0; i < 3 && n.parentElement; i++) {
      n = n.parentElement;
      const inputs = n.querySelectorAll('input[type=radio],input[type=checkbox],[role=radio],[role=checkbox]');
      if (inputs.length === 1 && n.innerText.trim()) { t = n; break; }
    }
  }
  return t ? fn(t, false).texto : (el.innerText || el.value || '');
}"""

# Datos de un <select>: opciones y el texto de su "fila" (ancestro más alto que solo contiene ese select).
_JS_HUECO = r"""(sel, selPregunta) => {
  const fn = """ + _JS_TEXTO_FN + r""";
  const ops = Array.from(sel.options).map((o, i) => ({ texto: o.text.trim(), valor: o.value, ind: i, dis: o.disabled }));
  let n = sel.parentElement, fila = null;
  while (n && !n.matches(selPregunta) && n.tagName !== 'BODY' && n.querySelectorAll('select').length === 1) { fila = n; n = n.parentElement; }
  const ctx = fila ? fn(fila, false).texto.replace(/\s+/g, ' ').trim() : '';
  return { ops: ops, ctx: ctx.length <= 400 ? ctx : '' };
}"""


def _limpiar(t: str) -> str:
    return re.sub(r"\s+", " ", t or "").strip()


def clasificar(pregunta: Locator) -> str | None:
    """Devuelve 'desplegable', 'multiple', 'unica', 'abierta' o None si no se reconoce."""
    # El orden importa: listas, checkbox y radio tienen prioridad sobre texto libre.
    if pregunta.locator(SELECTOR_DESPLEGABLE).count() > 0:
        return "desplegable"
    if pregunta.locator(SELECTOR_CHECKBOX).count() > 0:
        return "multiple"
    if pregunta.locator(SELECTOR_RADIO).count() > 0:
        return "unica"
    if pregunta.locator(SELECTOR_ABIERTA).count() > 0:
        return "abierta"
    return None


def _texto_lineas(t: str) -> str:
    """Conserva los saltos de línea (útil para incisos a), b), c)...) pero sin espacios sobrantes."""
    return "\n".join(l for l in (re.sub(r"[ \t\u00a0]+", " ", x).strip() for x in (t or "").splitlines()) if l)


def capturar_imagenes(pregunta: Locator, ids: list[str], maximo: int = 4) -> list[bytes]:
    """Captura (PNG) las imágenes/gráficas marcadas por _JS_TEXTO_FN, tal como se ven en pantalla."""
    out: list[bytes] = []
    for id_ in ids[:maximo]:
        try:
            png = pregunta.locator(f'[data-af-img="{id_}"]').first.screenshot(timeout=TIMEOUT_ACCION_MS)
            if len(png) <= 4_500_000:
                out.append(png)
        except Exception as e:
            print(f"     (no se pudo capturar una imagen: {str(e).splitlines()[0][:70]})")
    return out


def extraer_contenido(pregunta: Locator, huecos: bool = False,
                      solo_enunciado: bool = True) -> tuple[str, list[bytes]]:
    """Enunciado (fórmulas en LaTeX, huecos como [[n]]) e imágenes de la pregunta.

    solo_enunciado=True lee únicamente SELECTOR_ENUNCIADO (para no mezclar el texto de las opciones);
    con False lee todo el bloque (preguntas abiertas: así no se pierden incisos fuera del enunciado)."""
    try:
        cont = pregunta.evaluate(_JS_CONTENIDO, huecos)
        texto, ids = cont["texto"], cont["ids"]
        if not huecos and solo_enunciado:  # si hay un elemento de enunciado, el texto sale solo de ahí
            el = pregunta.locator(SELECTOR_ENUNCIADO).first
            if el.count() > 0:
                texto = el.evaluate(_JS_CONTENIDO, False)["texto"]
        return _texto_lineas(texto), capturar_imagenes(pregunta, ids)
    except Exception:
        return _texto_lineas(pregunta.inner_text(timeout=TIMEOUT_ACCION_MS)), []  # respaldo


_PLACEHOLDER = re.compile(r"^[\s\[\(\-–—_.]*(elegir|selecciona|seleccione|escoge|choose|select)?[\s\]\)\-–—_.…]*$", re.I)


def extraer_huecos(pregunta: Locator) -> list[dict]:
    """Por cada <select>: opciones reales (sin el '[ Elegir ]'), su posición real y el texto de su fila."""
    selects = pregunta.locator(SELECTOR_DESPLEGABLE)
    huecos = []
    for i in range(selects.count()):
        sel = selects.nth(i)
        datos = sel.evaluate(_JS_HUECO, SELECTOR_PREGUNTA)
        reales = [o for o in datos["ops"]
                  if o["valor"] != "" and not o["dis"] and not _PLACEHOLDER.match(o["texto"] or "")]
        if reales:
            huecos.append({"campo": sel, "textos": [_limpiar(o["texto"]) for o in reales],
                           "indices": [o["ind"] for o in reales], "contexto": datos["ctx"]})
    return huecos


def es_campo_corto(pregunta: Locator) -> bool:
    """True si la respuesta va en un campo de una sola línea (no en un textarea/editor)."""
    largos = pregunta.locator("textarea:visible, [contenteditable='true']:visible, " + SELECTOR_EDITOR_IFRAME)
    cortos = pregunta.locator("input[type='text']:visible, input[type='number']:visible, input:not([type]):visible")
    return largos.count() == 0 and cortos.count() > 0


def escribir_abierta(pregunta: Locator, texto: str) -> None:
    """Escribe en un campo de texto normal o, si no hay, en un editor enriquecido (contenteditable / iframe)."""
    normales = pregunta.locator("textarea:visible, input[type='text']:visible, input[type='number']:visible, "
                                "input:not([type]):visible")
    if normales.count() > 0:
        normales.first.fill(texto, timeout=TIMEOUT_ACCION_MS)
        return
    js = ("(el, t) => { el.focus(); el.innerHTML = t.split('\\n').map(l => '<p>' + l.replace(/&/g, '&amp;')"
          ".replace(/</g, '&lt;') + '</p>').join(''); "
          "el.dispatchEvent(new Event('input', {bubbles: true})); el.dispatchEvent(new Event('change', {bubbles: true})); }")
    editor = pregunta.locator("[contenteditable='true']:visible")
    if editor.count() > 0:
        editor.first.evaluate(js, texto)
        return
    if pregunta.locator(SELECTOR_EDITOR_IFRAME).count() > 0:  # TinyMCE (Canvas): el editor vive en un iframe
        pregunta.frame_locator(SELECTOR_EDITOR_IFRAME).first.locator("body").evaluate(js, texto)
        return
    pregunta.locator("textarea").first.evaluate(  # último recurso: textarea oculto
        "(el, t) => { el.value = t; el.dispatchEvent(new Event('input', {bubbles: true})); "
        "el.dispatchEvent(new Event('change', {bubbles: true})); }", texto)


def extraer_opciones(pregunta: Locator, selector_campo: str) -> tuple[list[str], list[Locator]]:
    """Devuelve (textos, campos) de las opciones, en el mismo orden."""
    campos = pregunta.locator(selector_campo)
    textos, locs = [], []
    for i in range(campos.count()):
        campo = campos.nth(i)
        try:
            texto = _limpiar(campo.evaluate(_JS_ETIQUETA))
        except Exception:
            texto = ""
        textos.append(texto or f"Opción {i + 1}")
        locs.append(campo)
    return textos, locs


def marcar(campo: Locator) -> None:
    """Marca un radio/checkbox; si el input está oculto por CSS, hace clic en su etiqueta."""
    try:
        campo.check(timeout=TIMEOUT_ACCION_MS)
        return
    except Exception:
        pass
    try:  # inputs ocultos con estilo personalizado
        campo.locator("xpath=ancestor-or-self::label[1]").first.click(timeout=TIMEOUT_ACCION_MS)
        return
    except Exception:
        pass
    try:  # roles ARIA (div role="radio")
        campo.click(timeout=TIMEOUT_ACCION_MS)
    except Exception:
        campo.check(force=True, timeout=TIMEOUT_ACCION_MS)


def responder_pregunta(pregunta: Locator, n: int, backend: str) -> None:
    tipo = clasificar(pregunta)
    if tipo is None:
        print(f"[{n}] Tipo no reconocido; se omite.")
        return

    enunciado, imagenes = extraer_contenido(pregunta, huecos=(tipo == "desplegable"),
                                            solo_enunciado=(tipo in ("unica", "multiple")))
    textos, campos, huecos = [], [], []
    if tipo == "unica":
        textos, campos = extraer_opciones(pregunta, SELECTOR_RADIO)
    elif tipo == "multiple":
        textos, campos = extraer_opciones(pregunta, SELECTOR_CHECKBOX)
    elif tipo == "desplegable":
        huecos = extraer_huecos(pregunta)
        if not huecos:
            print(f"[{n}] (DESPLEGABLE) sin opciones reales; se omite.")
            return

    dificil = clasificar_dificultad(enunciado, textos, backend)
    modelo = MODELO_DIFICIL if dificil else MODELO_FACIL
    etiqueta = ("DIFÍCIL" if dificil else "fácil") if ENRUTAR else "modelo único"
    extra = f" +{len(imagenes)} img" if imagenes else ""
    resumen = enunciado.replace("\n", " ")[:80]
    print(f"[{n}] ({tipo.upper()} | {etiqueta} | {modelo}{extra}) {resumen}")

    if tipo == "abierta":
        respuesta = pedir_abierta(enunciado, backend, modelo, dificil, imagenes, corta=es_campo_corto(pregunta))
        escribir_abierta(pregunta, respuesta)
        print(f"     -> {respuesta[:90]}")

    elif tipo == "unica":
        idx = pedir_unica(enunciado, textos, backend, modelo, imagenes)
        if idx is None:
            print("     -> Claude no devolvió una opción válida; se omite.")
            return
        marcar(campos[idx])
        print(f"     -> {textos[idx][:90]}")

    elif tipo == "multiple":
        indices = pedir_multiple(enunciado, textos, backend, modelo, imagenes)
        if not indices:
            print("     -> Claude no devolvió opciones válidas; se omite.")
            return
        for idx in indices:
            marcar(campos[idx])
        print(f"     -> {[textos[i][:40] for i in indices]}")

    elif tipo == "desplegable":
        elegidas = pedir_desplegables(enunciado, huecos, backend, modelo, imagenes)
        resumen_sel = []
        for h, e in zip(huecos, elegidas):
            if e is None:
                resumen_sel.append("?")
                continue
            try:
                h["campo"].select_option(index=h["indices"][e], timeout=TIMEOUT_ACCION_MS)
                resumen_sel.append(h["textos"][e][:20])
            except Exception as ex:
                resumen_sel.append("ERROR")
                print(f"     (no se pudo elegir en un desplegable: {str(ex).splitlines()[0][:60]})")
        print(f"     -> {resumen_sel}")


def procesar_pagina(page: Page, backend: str) -> None:
    preguntas = page.locator(SELECTOR_PREGUNTA)
    total = preguntas.count()
    if total == 0:
        print(f"No se encontró ninguna pregunta con el selector '{SELECTOR_PREGUNTA}'.")
        print("Ajusta SELECTOR_PREGUNTA al inicio del script.")
        return
    print(f"\nSe encontraron {total} preguntas.\n")

    for i in range(total):
        try:
            responder_pregunta(preguntas.nth(i), i + 1, backend)
        except Exception as e:  # un fallo en una pregunta no detiene las demás
            print(f"[{i + 1}] ERROR: {type(e).__name__}: {str(e)[:150]}")

    if SELECTOR_BOTON_ENVIAR:
        try:
            page.locator(SELECTOR_BOTON_ENVIAR).first.click(timeout=TIMEOUT_ACCION_MS)
            print("Formulario enviado.")
        except Exception as e:
            print(f"No se pudo enviar: {e}")
    else:
        print("\nListo. No se envió el formulario: revisa las respuestas y envíalo tú.")


# ════════════════════════════════════════════════════════════════════════════
# Formulario de demostración local (para probar sin ningún sitio real)
# ════════════════════════════════════════════════════════════════════════════
DEMO_HTML = """<!doctype html><meta charset=utf-8><title>Demo</title>
<div class=pregunta><p class=enunciado>1. ¿Cuál es la capital de Francia?</p>
  <textarea name=q1></textarea></div>
<div class=pregunta><p class=enunciado>2. ¿Cuál es el planeta más grande?</p>
  <div class=opcion><label><input type=radio name=q2> Marte</label></div>
  <div class=opcion><label><input type=radio name=q2> Júpiter</label></div>
  <div class=opcion><label><input type=radio name=q2> Venus</label></div></div>
<div class=pregunta><p class=enunciado>3. Selecciona los números primos</p>
  <div class=opcion><input type=checkbox id=a><label for=a>2</label></div>
  <div class=opcion><input type=checkbox id=b><label for=b>4</label></div>
  <div class=opcion><input type=checkbox id=c><label for=c>5</label></div>
  <div class=opcion><input type=checkbox id=d><label for=d>9</label></div></div>
<button type=submit>Enviar</button>"""


def crear_demo() -> str:
    ruta = Path(tempfile.mkdtemp()) / "demo.html"
    ruta.write_text(DEMO_HTML, encoding="utf-8")
    return ruta.as_uri()


# Demo que VERIFICA el uso del material: preguntas cuyas respuestas salen de los .txt de
# Teoría de Juegos / Economía Administrativa (valores de las matrices, láminas, etc.).
# "ok" = palabras que debe contener la respuesta abierta, o los índices (0-based) correctos.
DEMO_MATERIAL = [
    {"tipo": "abierta", "dif": "fácil",
     "q": "¿Cómo se llama el curso al que pertenecen las presentaciones de Teoría de Juegos con información incompleta?",
     "ok": ["decisiones estrat"]},
    {"tipo": "unica", "dif": "fácil",
     "q": "En el dilema del prisionero con información incompleta (presentación TJ-9), ¿qué pago recibe el "
          "sospechoso 1 si ambos hablan y el sospechoso 2 es de tipo Compa?",
     "ops": ["−6", "−9", "−3", "2"], "ok": [1]},
    {"tipo": "unica", "dif": "difícil",
     "q": "En el ejemplo de selección adversa de la presentación TJ-III, si el comprador ofrece p = 20, "
          "¿cuánto vale E[v2 | venta]?",
     "ops": ["24", "20", "19", "14"], "ok": [2]},
    {"tipo": "unica", "dif": "fácil",
     "q": "En la batalla de los sexos bayesiana de TJ-9, si Cris elige Cine con certeza y π ≥ 1/4, ¿qué perfil "
          "(λ, µG, µS) es un equilibrio bayesiano de Nash en estrategias puras?",
     "ops": ["(0, 0, 1)", "(3/4, 1/4, 1/4)", "(1, 1, 0)", "(1, 0, 1)"], "ok": [2]},
    {"tipo": "abierta", "dif": "difícil",
     "q": "En el dilema del prisionero con información incompleta de la presentación TJ-9, el sospechoso 1 cree "
          "que el sospechoso 2 es Judas con probabilidad ρ. Calcula el valor de ρ que hace indiferente al "
          "sospechoso 1 entre Hablar y Callar. Muestra el desarrollo.",
     "ok": ["5/8", "0.625", "0,625"]},
    {"tipo": "abierta", "dif": "fácil",
     "q": "En el caso FríoTech México del Apunte 2, ¿cuántos controladores electrónicos necesita la empresa al año?",
     "ok": ["15,000", "15000", "15 000", "15.000"]},
    {"tipo": "unica", "dif": "fácil",
     "q": "En el ejemplo de integración atemperada (Make AND Buy) del Apunte 5, con una demanda total de "
          "100,000 unidades, ¿cuántas se compran al proveedor externo?",
     "ops": ["40,000", "70,000", "60,000", "100,000"], "ok": [2]},
    {"tipo": "multiple", "dif": "fácil",
     "q": "Según el Apunte 1, ¿qué dimensiones de la transacción destaca Williamson?",
     "ops": ["Tamaño de la empresa", "Frecuencia", "Reputación del proveedor",
             "Especificidad de activos", "Incertidumbre"], "ok": [1, 3, 4]},
    {"tipo": "multiple", "dif": "fácil",
     "q": "Según las láminas de TJ-IV (dos jurados, regla de unanimidad), ¿qué afirmaciones son correctas?",
     "ops": ["(CA, CA) es un equilibrio bayesiano de Nash.",
             "Un jurado solo es pivotal si el otro jurado vota C.",
             "Sin señales y con q > 1/2, votar C es una estrategia débilmente dominante.",
             "La persona acusada es condenada si al menos un jurado vota C."], "ok": [1, 2]},
]


def _svg_url(svg: str) -> str:
    return "data:image/svg+xml;utf8," + quote(svg)


def _svg_grafica(etiqueta_b: str) -> str:
    """Isocosto (recta roja) tangente a una isocuanta (curva negra); a y b son puntos de los ejes."""
    return (
        "<svg xmlns='http://www.w3.org/2000/svg' width='480' height='320' font-family='serif' font-size='26'>"
        "<rect width='480' height='320' fill='white'/>"
        "<line x1='60' y1='20' x2='60' y2='270' stroke='black' stroke-width='3'/>"
        "<line x1='40' y1='270' x2='450' y2='270' stroke='black' stroke-width='3'/>"
        "<text x='20' y='40'>K</text><text x='440' y='300'>L</text>"
        "<line x1='60' y1='70' x2='300' y2='270' stroke='#e5483f' stroke-width='3'/>"
        "<path d='M110 40 Q150 190 300 235 T440 262' fill='none' stroke='black' stroke-width='3'/>"
        "<text x='30' y='78'>a</text>"
        f"<text x='290' y='305'>{etiqueta_b}</text></svg>")


def crear_demo_formatos() -> str:
    """Formulario local con los formatos típicos de Canvas: desplegables V/F, ejercicio con fórmula
    (imagen de ecuación) y editor enriquecido en iframe, y preguntas con gráficas."""
    ecuacion = ("<img class='equation_image' alt='LaTeX: U(X_A,X_B)=\\sqrt{X_AX_B+400}' "
                "data-equation-content='U(X_A,X_B)=\\sqrt{X_AX_B+400}' width='220' height='30' "
                f"src=\"{_svg_url('<svg xmlns=\'http://www.w3.org/2000/svg\' width=\'220\' height=\'30\'/>')}\">")
    ec_iso = ("<img class='equation_image' alt='LaTeX: C=wL+rK' data-equation-content='C=wL+rK' width='110' height='24' "
              f"src=\"{_svg_url('<svg xmlns=\'http://www.w3.org/2000/svg\' width=\'110\' height=\'24\'/>')}\">")
    filas = "".join(
        f"<div class=fila><span>{txt}</span><select><option value=''>[ Elegir ]</option>"
        "<option>Verdadero</option><option>Falso</option></select></div>"
        for txt, _ in DEMO_VF)
    html = f"""<!doctype html><meta charset=utf-8><title>Demo formatos</title>
<style>.fila{{display:flex;justify-content:space-between;gap:20px;padding:8px;border-top:1px solid #ccc}}</style>
<div class=pregunta><p class=enunciado>Verdadero o falso</p>{filas}</div>
<div class=pregunta><div class=enunciado><p>La función de utilidad del Pato Lucas es:</p><p>{ecuacion}</p>
<p>a) Obtén las condiciones de primer orden. (4 puntos)</p><p>b) Obtén la tasa marginal de sustitución. (4 puntos)</p>
<p>c) Obtén la función de demanda del bien A. (4 puntos)</p><p>d) Obtén la demanda del bien B. (4 puntos)</p>
<p>e) Determina si el bien A y el bien B son sustitutos o complementos. (4 puntos)</p></div>
<iframe id=q2_ifr style='width:600px;height:140px' srcdoc="<body contenteditable=true></body>"></iframe></div>
<div class=pregunta><div class=enunciado><p>La recta roja representa la curva de isocosto: {ec_iso}</p>
<img src="{_svg_url(_svg_grafica('b'))}" width=480 height=320>
<p>Si C=600, w=4 y r=3, determina el valor del punto a.</p></div><input type=text></div>
<div class=pregunta><div class=enunciado><p>Según la gráfica, ¿cuánto vale el punto b?</p>
<img src="{_svg_url(_svg_grafica('b = 150'))}" width=480 height=320></div><input type=text></div>"""
    ruta = Path(tempfile.mkdtemp()) / "demo_formatos.html"
    ruta.write_text(html, encoding="utf-8")
    return ruta.as_uri()


# Verdadero/falso del demo de formatos (respuestas según la clave del examen de ejemplo).
DEMO_VF = [
    ("La función de costos es igual que la recta de isocosto", "Falso"),
    ("Si tengo pérdidas, siempre debo dejar de producir.", "Falso"),
    ("Solamente tengo costos fijos en el corto plazo.", "Verdadero"),
    ("La producción óptima se encuentra en el punto de intersección entre costos e ingresos.", "Falso"),
    ("Una empresa en un mercado competitivo que no tenga costos fijos y que tenga un costo marginal constante y "
     "positivo, producirá hasta que su ingreso marginal sea cero.", "Falso"),
    ("La derivada de la función de beneficios debe ser positiva para garantizar una producción eficiente y óptima.",
     "Falso"),
    ("En el corto plazo, hay al menos un factor de producción fijo.", "Verdadero"),
    ("Un rendimiento decreciente de escala implica que obtengo pérdidas si aumento la producción.", "Falso"),
]


def evaluar_demo_formatos(page: Page) -> None:
    preguntas = page.locator(SELECTOR_PREGUNTA)
    print("\n" + "=" * 60 + "\n RESULTADO DE LA DEMO DE FORMATOS\n" + "=" * 60)
    # P1: desplegables
    selects = preguntas.nth(0).locator("select")
    ok_vf = 0
    for i, (txt, esperado) in enumerate(DEMO_VF):
        elegido = selects.nth(i).evaluate("s => s.options[s.selectedIndex].text")
        ok_vf += elegido == esperado
        print(f"   {'✔' if elegido == esperado else '✘'} {txt[:55]:<55} -> {elegido} (esperado {esperado})")
    print(f" P1 desplegables V/F: {ok_vf}/{len(DEMO_VF)}")
    # P2: ejercicio con fórmula en editor iframe
    cuerpo = preguntas.nth(1).frame_locator("iframe").locator("body").inner_text().lower()
    checks = {"escribió en el editor": len(cuerpo) > 40,
              "respondió los incisos a) a e)": all(f"{c})" in cuerpo for c in "abcde"),
              "vio la fórmula (aparece 400)": "400" in cuerpo}
    print(f" P2 Pato Lucas: " + " | ".join(f"{'✔' if v else '✘'} {k}" for k, v in checks.items()))
    # P3 / P4: gráficas
    v3 = preguntas.nth(2).locator("input").input_value()
    v4 = preguntas.nth(3).locator("input").input_value()
    print(f" P3 isocosto (a = C/r = 200): {'✔' if '200' in v3 else '✘'} -> {v3[:60]!r}")
    print(f" P4 valor leído solo de la imagen (b = 150): {'✔' if '150' in v4 else '✘'} -> {v4[:60]!r}")


def crear_demo_material() -> str:
    bloques = []
    for n, q in enumerate(DEMO_MATERIAL, 1):
        if q["tipo"] == "abierta":
            campo = "<textarea rows=4 cols=60></textarea>"
        else:
            kind = "radio" if q["tipo"] == "unica" else "checkbox"
            campo = "".join(
                f"<div class=opcion><label><input type={kind} name=q{n}> {op}</label></div>"
                for op in q["ops"])
        bloques.append(f"<div class=pregunta><p class=enunciado>{n}. {q['q']}</p>{campo}</div>")
    html = "<!doctype html><meta charset=utf-8><title>Demo material</title>" + "\n".join(bloques)
    ruta = Path(tempfile.mkdtemp()) / "demo_material.html"
    ruta.write_text(html, encoding="utf-8")
    return ruta.as_uri()


def evaluar_demo_material(page: Page) -> None:
    """Califica las respuestas de la demo contra las respuestas esperadas."""
    preguntas = page.locator(SELECTOR_PREGUNTA)
    aciertos = 0
    print("\n" + "=" * 60 + "\n RESULTADO DE LA DEMO (¿usó tu material?)\n" + "=" * 60)
    for i, q in enumerate(DEMO_MATERIAL):
        bloque = preguntas.nth(i)
        if q["tipo"] == "abierta":
            valor = bloque.locator("textarea").first.input_value().lower()
            ok = any(k in valor for k in q["ok"])
            esperado = " / ".join(q["ok"])
        else:
            campos = bloque.locator("input")
            marcados = [j for j in range(campos.count()) if campos.nth(j).is_checked()]
            ok = marcados == q["ok"]
            esperado = ", ".join(q["ops"][j] for j in q["ok"])
        aciertos += ok
        print(f" {'✔' if ok else '✘'} P{i + 1} ({q['tipo']}, {q['dif']}) esperado: {esperado}")
    print(f"\n Aciertos: {aciertos}/{len(DEMO_MATERIAL)}")
    if not CONTEXTO_MATERIAL:
        print(" (Corriste SIN --material: esto es la línea base con conocimiento general.)")


# ════════════════════════════════════════════════════════════════════════════
# main
# ════════════════════════════════════════════════════════════════════════════
def elegir_backend(pedido: str) -> str:
    if pedido != "auto":
        return pedido
    if os.getenv("ANTHROPIC_API_KEY"):
        return "api"
    if shutil.which("claude"):
        return "cli"
    sys.exit("Define ANTHROPIC_API_KEY o instala el CLI `claude` (o usa --backend mock).")


def main() -> None:
    global MODELO_FACIL, MODELO_DIFICIL, ENRUTAR, CONTEXTO_MATERIAL
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    ap.add_argument("--url", help="URL inicial a abrir (opcional)")
    ap.add_argument("--backend", choices=["auto", "api", "cli", "mock"], default="auto")
    ap.add_argument("--material", help="archivo o carpeta (.pptx/.pdf/.txt/.md) que Claude usará como base")
    ap.add_argument("--model", help="un solo modelo para TODAS las preguntas (desactiva el enrutamiento)")
    ap.add_argument("--model-facil", help=f"modelo para preguntas fáciles (activa el enrutamiento si difiere del difícil; por defecto {MODELO_FACIL})")
    ap.add_argument("--model-dificil", help=f"modelo para preguntas difíciles (por defecto {MODELO_DIFICIL})")
    ap.add_argument("--demo-formatos", action="store_true",
                    help="formulario local con desplegables V/F, fórmulas, editor enriquecido y gráficas")
    ap.add_argument("--demo-material", action="store_true",
                    help="formulario local con preguntas que solo se contestan con tu material, y lo califica")
    ap.add_argument("--demo", action="store_true", help="usa un formulario local de ejemplo")
    ap.add_argument("--headless", action="store_true", help="sin ventana (solo pruebas)")
    ap.add_argument("--auto-start", action="store_true", help="no espera ENTER (solo pruebas)")
    args = ap.parse_args()

    if args.model:
        MODELO_FACIL = MODELO_DIFICIL = args.model
        ENRUTAR = False
    if args.model_facil:
        MODELO_FACIL = args.model_facil
    if args.model_dificil:
        MODELO_DIFICIL = args.model_dificil
    ENRUTAR = MODELO_FACIL != MODELO_DIFICIL
    backend = elegir_backend(args.backend)
    if args.demo_formatos:
        url = crear_demo_formatos()
    else:
        url = crear_demo_material() if args.demo_material else (crear_demo() if args.demo else args.url)
    modelos = MODELO_FACIL if not ENRUTAR else f"fácil={MODELO_FACIL}, difícil={MODELO_DIFICIL}"
    print(f"Backend de Claude: {backend} | modelos: {modelos}")
    if args.material:
        print("Cargando material de consulta...")
        CONTEXTO_MATERIAL = cargar_material(args.material)

    with sync_playwright() as p:
        try:
            browser = p.chromium.launch(headless=args.headless)
        except Exception:
            # Respaldo: usar un Chromium ya instalado en el sistema (variable opcional).
            ruta = os.getenv("CHROMIUM_PATH")
            if not ruta:
                raise
            browser = p.chromium.launch(headless=args.headless, executable_path=ruta)
        page = browser.new_page()
        if url:
            try:
                page.goto(url)
            except Exception as e:  # URL mal escrita o sin conexión: no abortar, navegas tú a mano
                print(f"AVISO: no se pudo abrir '{url}' ({str(e).splitlines()[0][:80]}).")
                print("Navega manualmente en el navegador a la página correcta.")

        if not args.auto_start:
            esperar_activador()
        procesar_pagina(page, backend)
        if args.demo_material:
            evaluar_demo_material(page)
        if args.demo_formatos:
            evaluar_demo_formatos(page)
        imprimir_consumo()

        if not args.headless:
            input("\nPresiona ENTER para cerrar el navegador... ")
        else:
            for etiqueta, selector in [("textarea", "textarea"), ("radio", "input[type=radio]:checked"),
                                       ("checkbox", "input[type=checkbox]:checked")]:
                print(f"[verificación] {etiqueta}:", page.locator(selector).count())
        browser.close()


if __name__ == "__main__":
    main()
