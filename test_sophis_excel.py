import datetime as dt

import pytest

import sophis_excel as xl


def row(date, value, code="SNDK"):
    return ("SNDK US Equity", f"SNDK US {date} C45", f"SNDK US {date} P45", date, value, code, "A", 100.0, "M", "N", 0, 0)


HEADER = ("Equity", "Call", "Put", "Date", "Value", "Code", "L1", "V", "L2", "L3", "V1", "V2")


def test_first_value_is_minimum_last_is_maximum_per_date():
    rows = [HEADER, row("01/15/27", 45.0), row("01/15/27", 50.0), row("01/15/27", 100.0), row("02/19/27", 30.0),
            row("02/19/27", 70.0)]
    plan, warnings = xl.extract(rows)
    assert list(plan) == ["SNDK"]
    got = [(d.date, d.minimum, d.maximum, d.rows) for d in plan["SNDK"]]
    assert got == [(dt.date(2027, 1, 15), 45.0, 100.0, 3), (dt.date(2027, 2, 19), 30.0, 70.0, 2)]
    assert warnings == []


def test_it_is_the_order_of_the_file_not_the_size():
    plan, warnings = xl.extract([row("01/15/27", 90.0), row("01/15/27", 40.0)])
    d = plan["SNDK"][0]
    assert (d.minimum, d.maximum) == (90.0, 40.0)
    assert any("bigger than" in w for w in warnings)


def test_several_codes_each_get_their_own_dates():
    rows = [row("01/15/27", 45.0, "SNDK"), row("01/15/27", 60.0, "SNDK"), row("01/15/27", 5.0, "AAPL"),
            row("03/19/27", 9.0, "AAPL"), row("03/19/27", 12.0, "AAPL")]
    plan, _ = xl.extract(rows)
    assert set(plan) == {"SNDK", "AAPL"}
    assert [(d.minimum, d.maximum) for d in plan["AAPL"]] == [(5.0, 5.0), (9.0, 12.0)]


def test_dates_are_sorted_and_a_single_value_is_both_ends_with_a_warning():
    plan, warnings = xl.extract([row("03/19/27", 9.0), row("01/15/27", 45.0)])
    assert [d.date for d in plan["SNDK"]] == [dt.date(2027, 1, 15), dt.date(2027, 3, 19)]
    assert sum("only one value" in w for w in warnings) == 2


def test_real_date_cells_and_text_numbers():
    plan, _ = xl.extract([("", "", "", dt.datetime(2027, 1, 15), "45.00", "SNDK"),
                          ("", "", "", dt.datetime(2027, 1, 15), "1,100.50", "SNDK")])
    d = plan["SNDK"][0]
    assert (d.minimum, d.maximum) == (45.0, 1100.5)


def test_unusable_rows_are_reported_and_headers_skipped():
    plan, warnings = xl.extract([HEADER, row("01/15/27", None), row("01/15/27", 45.0), (None,) * 12])
    assert len(plan["SNDK"]) == 1 and any("row 2" in w for w in warnings)


def test_number_format_and_entries():
    assert [xl.format_number(v) for v in (45.0, 45.5, 0.0, 100.25)] == ["45", "45.5", "0", "100.25"]
    plan, _ = xl.extract([row("01/15/27", 45.0), row("01/15/27", 100.0)])
    assert xl.entries_for(plan["SNDK"]) == [("15 Jan 2027", "45", "100")]


def test_reads_a_real_xlsx_and_a_csv(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(HEADER)
    for r in (row("01/15/27", 45.0), row("01/15/27", 100.0)):
        ws.append(r)
    ws.append(("", "", "", dt.datetime(2027, 2, 19), 30, "SNDK"))
    path = tmp_path / "strikes.xlsx"
    wb.save(path)
    plan, _ = xl.read_plan(path)
    assert [(d.date, d.minimum, d.maximum) for d in plan["SNDK"]] == [
        (dt.date(2027, 1, 15), 45.0, 100.0), (dt.date(2027, 2, 19), 30.0, 30.0)]
    csv_path = tmp_path / "strikes.csv"
    csv_path.write_text("a,b,c,d,e,f\n,,,01/15/27,45.00,SNDK\n,,,01/15/27,100.00,SNDK\n")
    assert xl.read_plan(csv_path)[0]["SNDK"][0].maximum == 100.0


def test_clear_errors(tmp_path):
    with pytest.raises(xl.ExcelError, match="does not exist"):
        xl.read_plan(tmp_path / "nope.xlsx")
    bad = tmp_path / "x.csv"
    bad.write_text("nothing,useful\n")
    with pytest.raises(xl.ExcelError, match="No data rows"):
        xl.read_plan(bad)
    old = tmp_path / "old.xls"
    old.write_text("x")
    with pytest.raises(xl.ExcelError, match="No data rows"):        # plain text with an .xls name: read as text
        xl.read_plan(old)
    html = tmp_path / "export.xls"
    html.write_text("<html><table></table></html>")
    with pytest.raises(xl.ExcelError, match="No data rows"):        # an HTML page without data rows
        xl.read_plan(html)


def _args(path, **kw):
    import argparse
    return argparse.Namespace(excel=str(path), sheet=None, code=None, **kw)


def test_fill_takes_the_entries_from_the_excel(tmp_path, capsys):
    import sophis_auto as sa
    f = tmp_path / "s.csv"
    f.write_text(",,,01/15/27,45.00,SNDK\n,,,01/15/27,100.00,SNDK\n,,,02/19/27,30.00,SNDK\n")
    triples = sa._excel_triples(_args(f))
    assert triples == [("15 Jan 2027", "45", "100"), ("19 Feb 2027", "30", "30")]
    assert sa.parse_entry(*triples[0]).date == dt.date(2027, 1, 15)


def test_fill_with_several_codes_needs_a_choice(tmp_path):
    import sophis_auto as sa
    f = tmp_path / "s.csv"
    f.write_text(",,,01/15/27,45.00,SNDK\n,,,01/15/27,5.00,AAPL\n")
    with pytest.raises(sa.SophisError, match="choose one with --code"):
        sa._excel_triples(_args(f))
    a = _args(f)
    a.code = "aapl"
    assert sa._excel_triples(a) == [("15 Jan 2027", "5", "5")]
    a.code = "XXX"
    with pytest.raises(sa.SophisError, match="not in the Excel"):
        sa._excel_triples(a)


def test_reads_an_old_xls_with_text_numbers_and_real_dates(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Data")
    for c, v in enumerate(HEADER):
        ws.write(0, c, v)
    date_style = xlwt.easyxf(num_format_str="MM/DD/YY")
    for r, (date, value) in enumerate([("01/15/27", 45.0), ("01/15/27", 100.0)], start=1):
        ws.write(r, 3, date)                                   # text date
        ws.write(r, 4, value)
        ws.write(r, 5, "SNDK")
    ws.write(3, 3, dt.datetime(2027, 2, 19), date_style)       # real date cell
    ws.write(3, 4, "30.00")                                    # number stored as text
    ws.write(3, 5, "SNDK")
    path = tmp_path / "strikes.xls"
    wb.save(str(path))
    plan, _ = xl.read_plan(path)
    assert [(d.date, d.minimum, d.maximum) for d in plan["SNDK"]] == [
        (dt.date(2027, 1, 15), 45.0, 100.0), (dt.date(2027, 2, 19), 30.0, 30.0)]
    assert xl.read_plan(path, "Data")[0].keys() == {"SNDK"}
    with pytest.raises(xl.ExcelError, match="No sheet named"):
        xl.read_plan(path, "Other")


def test_a_semicolon_or_tab_text_export_and_an_xlsx_named_xls(tmp_path):
    tab = tmp_path / "t.txt"
    tab.write_text("a\tb\tc\td\te\tf\n\t\t\t01/15/27\t45.00\tSNDK\n\t\t\t01/15/27\t100.00\tSNDK\n")
    assert xl.read_plan(tab)[0]["SNDK"][0].maximum == 100.0
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(("", "", "", "01/15/27", 45.0, "SNDK"))
    path = tmp_path / "disguised.xls"
    wb.save(path)
    assert xl.read_plan(path)[0]["SNDK"][0].minimum == 45.0


def test_an_xls_that_is_really_an_html_table(tmp_path):
    path = tmp_path / "export.xls"
    path.write_text("<html><body><table><tr><th>Equity</th><th>Date</th></tr>"
                    "<tr><td>SNDK US Equity</td><td>x</td><td>y</td><td>01/15/27</td><td>45.00</td><td>SNDK</td></tr>"
                    "<tr><td></td><td></td><td></td><td>01/15/27</td><td>1,100.50</td><td>SNDK</td></tr>"
                    "</table></body></html>", encoding="utf-8")
    plan, _ = xl.read_plan(path)
    d = plan["SNDK"][0]
    assert (d.date, d.minimum, d.maximum) == (dt.date(2027, 1, 15), 45.0, 1100.5)


def test_an_xls_that_is_really_an_xml_spreadsheet(tmp_path):
    path = tmp_path / "export.xls"
    path.write_text('''<?xml version="1.0"?>
<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">
 <Worksheet ss:Name="Data"><Table>
  <Row><Cell><Data ss:Type="String">h</Data></Cell></Row>
  <Row><Cell ss:Index="4"><Data ss:Type="String">01/15/27</Data></Cell><Cell><Data ss:Type="Number">45</Data></Cell>
       <Cell><Data ss:Type="String">SNDK</Data></Cell></Row>
  <Row><Cell ss:Index="4"><Data ss:Type="String">01/15/27</Data></Cell><Cell><Data ss:Type="Number">100</Data></Cell>
       <Cell><Data ss:Type="String">SNDK</Data></Cell></Row>
 </Table></Worksheet>
</Workbook>''', encoding="utf-8")
    plan, _ = xl.read_plan(path)
    assert (plan["SNDK"][0].minimum, plan["SNDK"][0].maximum) == (45.0, 100.0)
    with pytest.raises(xl.ExcelError, match="No sheet named"):
        xl.read_plan(path, "Other")


def test_a_broken_xls_gives_a_clear_message(tmp_path):
    path = tmp_path / "broken.xls"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 40)
    with pytest.raises(xl.ExcelError, match="save it as .xlsx"):
        xl.read_plan(path)


def test_the_file_picker_falls_back_to_the_console_without_a_screen(monkeypatch):
    import builtins
    real_import = builtins.__import__
    def no_tk(name, *a, **k):
        if name.startswith("tkinter"):
            raise ImportError("no display")
        return real_import(name, *a, **k)
    monkeypatch.setattr(builtins, "__import__", no_tk)
    assert xl.pick_file(input_fn=lambda prompt: ' "C:\\data\\strikes.xls" ') == "C:\\data\\strikes.xls"
    assert xl.pick_file(input_fn=lambda prompt: "") is None


def test_the_file_picker_returns_what_the_window_chose(monkeypatch):
    import sys, types
    root = types.SimpleNamespace(withdraw=lambda: None, attributes=lambda *a: None, destroy=lambda: None)
    tk = types.ModuleType("tkinter")
    tk.Tk = lambda: root
    fd = types.ModuleType("tkinter.filedialog")
    seen = {}
    def ask(**kw):
        seen.update(kw)
        return "/data/strikes.xls"
    fd.askopenfilename = ask
    tk.filedialog = fd
    monkeypatch.setitem(sys.modules, "tkinter", tk)
    monkeypatch.setitem(sys.modules, "tkinter.filedialog", fd)
    assert xl.pick_file() == "/data/strikes.xls"
    assert any("*.xls" in pattern for _, pattern in seen["filetypes"])


def test_run_without_a_file_opens_the_picker(monkeypatch, tmp_path, capsys):
    import sophis_auto as sa
    f = tmp_path / "s.csv"
    f.write_text(",,,01/15/27,45.00,SNDK\n")
    monkeypatch.setattr(xl, "pick_file", lambda *a, **k: str(f))
    class Stop(Exception):
        pass
    def boom(args):
        raise Stop
    monkeypatch.setattr(sa, "_create", boom)
    monkeypatch.setattr("sys.argv", ["x", "run"])
    with pytest.raises(Stop):
        sa.main()
    out = capsys.readouterr().out
    assert "Search FIRST in Sophis: SNDK" in out and str(f) in out
    monkeypatch.setattr(xl, "pick_file", lambda *a, **k: None)
    with pytest.raises(SystemExit):
        sa.main()


def test_an_outdated_sophis_excel_is_reported_clearly(monkeypatch, capsys):
    import sophis_auto as sa
    monkeypatch.delattr(xl, "pick_file")
    with pytest.raises(SystemExit):
        sa._import_excel_module()
    out = capsys.readouterr().out
    assert "out of date" in out and "pick_file" in out and "git pull" in out
