"""Reads the Excel (.xlsx / .xls) or CSV with the option strikes and turns it into the dates and Minimum / Maximum values to type
in Sophis.

Columns (by position, the header row, if any, is skipped):
    A  Bloomberg equity       "SNDK US Equity"
    B  call code              "SNDK US 01/15/27 C45"
    C  put code               "SNDK US 01/15/27 P45"
    D  date                   01/15/27            <- filter
    E  value                  45.00               <- first one of a date = Minimum, last one = Maximum
    F  code                   "SNDK"              <- one plan per code
    G..L  not used

Usage:  python sophis_excel.py FILE.xlsx [--sheet NAME] [--code SNDK]
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

COL_DATE, COL_VALUE, COL_CODE = 3, 4, 5      # D, E, F (0-based)


class ExcelError(Exception):
    pass


@dataclass
class DateValues:
    date: dt.date
    minimum: float                # first value of the date, in file order
    maximum: float                # last value of the date, in file order
    rows: int                     # how many rows the date had
    excel_rows: list = field(default_factory=list)   # 1-based sheet rows, for messages

    @property
    def suspicious(self) -> bool:
        return self.minimum > self.maximum


def format_number(value: float) -> str:
    """45.0 -> '45', 45.5 -> '45.5' (what is typed in the Sophis cell)."""
    text = f"{value:.6f}".rstrip("0").rstrip(".")
    return text if text not in ("", "-0") else "0"


def parse_excel_date(value) -> dt.date | None:
    """Real date cells, or text like 01/15/27, 01/15/2027, 2027-01-15 (the file uses month/day/year)."""
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if value is None:
        return None
    text = str(value).strip()
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{2}|\d{4})", text)
    if m:
        month, day, year = int(m[1]), int(m[2]), int(m[3])
        year += 2000 if year < 100 else 0
        try:
            return dt.date(year, month, day)
        except ValueError:
            return None
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})(?:[ T].*)?", text)
    if m:
        try:
            return dt.date(int(m[1]), int(m[2]), int(m[3]))
        except ValueError:
            return None
    return None


def parse_excel_number(value) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def _sniff(path: Path) -> str:
    """The real kind of the file, from its first bytes (an .xls is often really an .xlsx or a text export)."""
    head = path.read_bytes()[:8]
    if head[:2] == b"PK":
        return "xlsx"
    if head == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "xls"
    if path.suffix.lower() in (".csv", ".txt", ".tsv"):
        return "text"
    start = path.read_bytes()[:512].lstrip().lower()
    if start.startswith((b"<html", b"<?xml", b"<table", b"<!doctype")):
        raise ExcelError(f"{path.name} is an HTML/XML export saved as {path.suffix}: open it in Excel and "
                         "save it as .xlsx or .xls first.")
    return "text"


def _load_xlsx(path: Path, sheet: str | None) -> list[tuple]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ExcelError("openpyxl is needed to read .xlsx files:  pip install openpyxl")
    import io
    wb = load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=True)   # (name-independent)
    try:
        if sheet is not None:
            if sheet not in wb.sheetnames:
                raise ExcelError(f"No sheet named {sheet!r}; the file has {wb.sheetnames}.")
            ws = wb[sheet]
        else:
            ws = wb.active
        return [tuple(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _load_xls(path: Path, sheet: str | None) -> list[tuple]:
    try:
        import xlrd
    except ImportError:
        raise ExcelError("xlrd is needed to read old .xls files:  pip install xlrd")
    book = xlrd.open_workbook(str(path))
    try:
        if sheet is not None:
            if sheet not in book.sheet_names():
                raise ExcelError(f"No sheet named {sheet!r}; the file has {book.sheet_names()}.")
            ws = book.sheet_by_name(sheet)
        else:
            ws = book.sheet_by_index(0)
        rows = []
        for r in range(ws.nrows):
            row = []
            for c in range(ws.ncols):
                cell = ws.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                    row.append(None)
                else:
                    row.append(cell.value)
            rows.append(tuple(row))
        return rows
    finally:
        book.release_resources()


def _load_text(path: Path) -> list[tuple]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    first = text.splitlines()[0] if text.strip() else ""
    delimiter = max(",;\t", key=first.count) if first else ","
    return [tuple(r) for r in csv.reader(text.splitlines(), delimiter=delimiter)]


def load_rows(path: str | Path, sheet: str | None = None) -> list[tuple]:
    """Every row of the file as a tuple of cell values (computed values, not formulas). The kind of file is
    decided by its content: .xlsx / .xlsm, old .xls, or text (.csv / .txt / .tsv)."""
    path = Path(path)
    if not path.exists():
        raise ExcelError(f"The file {path} does not exist.")
    kind = _sniff(path)
    if kind == "xlsx":
        return _load_xlsx(path, sheet)
    if kind == "xls":
        return _load_xls(path, sheet)
    return _load_text(path)


def extract(rows: list[tuple]) -> tuple[dict[str, list[DateValues]], list[str]]:
    """{code: [DateValues, ...]} with the dates in ascending order, plus warnings.

    Per code and date, the values of column E are taken in the order of the file: the first is the Minimum, the
    last is the Maximum. Rows that are not data (header, blank, totals) are skipped; a row with a date but a
    code or value that cannot be read is reported, not silently dropped."""
    groups: dict[str, dict[dt.date, list[tuple[int, float]]]] = {}
    warnings: list[str] = []
    for n, row in enumerate(rows, start=1):
        row = tuple(row) + (None,) * (COL_CODE + 1 - len(row))
        date = parse_excel_date(row[COL_DATE])
        if date is None:
            continue                                         # header / blank / not a data row
        code = str(row[COL_CODE]).strip() if row[COL_CODE] is not None else ""
        value = parse_excel_number(row[COL_VALUE])
        if not code or value is None:
            warnings.append(f"row {n}: date {date:%m/%d/%y} but code {row[COL_CODE]!r} / value {row[COL_VALUE]!r} "
                            "cannot be used; row skipped")
            continue
        groups.setdefault(code, {}).setdefault(date, []).append((n, value))
    plan: dict[str, list[DateValues]] = {}
    for code, by_date in groups.items():
        items = []
        for date in sorted(by_date):
            vals = by_date[date]
            dv = DateValues(date, vals[0][1], vals[-1][1], len(vals), [n for n, _ in vals])
            if dv.rows == 1:
                warnings.append(f"{code} {date:%m/%d/%y}: only one value ({format_number(dv.minimum)}), "
                                "used as both Minimum and Maximum")
            if dv.suspicious:
                warnings.append(f"{code} {date:%m/%d/%y}: first value {format_number(dv.minimum)} is bigger than "
                                f"the last {format_number(dv.maximum)}; check the order of the rows")
            items.append(dv)
        plan[code] = items
    return plan, warnings


def read_plan(path: str | Path, sheet: str | None = None) -> tuple[dict[str, list[DateValues]], list[str]]:
    plan, warnings = extract(load_rows(path, sheet))
    if not plan:
        raise ExcelError("No data rows found: column D must hold the dates (e.g. 01/15/27), column E the values "
                         "and column F the code.")
    return plan, warnings


def entries_for(items: list[DateValues]) -> list[tuple[str, str, str]]:
    """(date text, Minimum, Maximum) triples in the form `fill --entry` takes, e.g. ('15 Jan 2027', '45', '100')."""
    return [(f"{d.date:%d %b %Y}", format_number(d.minimum), format_number(d.maximum)) for d in items]


def describe(plan: dict[str, list[DateValues]]) -> str:
    lines = []
    for code, items in plan.items():
        lines.append(f"{code}: {len(items)} date(s)")
        for d in items:
            lines.append(f"   {d.date:%d %b %Y}   Minimum {format_number(d.minimum):>8}   "
                         f"Maximum {format_number(d.maximum):>8}   ({d.rows} row(s))")
    return "\n".join(lines)


def main(argv=None) -> None:
    p = argparse.ArgumentParser(description="Show the dates and Minimum / Maximum values found in the Excel.")
    p.add_argument("file")
    p.add_argument("--sheet", default=None, help="sheet name (default: the active sheet; the first one for .xls)")
    p.add_argument("--code", default=None, help="only this code")
    args = p.parse_args(argv)
    try:
        plan, warnings = read_plan(args.file, args.sheet)
        if args.code:
            plan = {k: v for k, v in plan.items() if k.upper() == args.code.upper()}
            if not plan:
                raise ExcelError(f"Code {args.code!r} is not in the file.")
    except ExcelError as e:
        print(f"ERROR: {e}")
        sys.exit(2)
    print(describe(plan))
    for w in warnings:
        print("WARNING:", w)


if __name__ == "__main__":
    main()
